"""Parse an uploaded resume (PDF or DOCX) into resume_content.yaml.

Used by the onboarding wizard: the user drops their resume, we extract the
plain text, ask an LLM to structure it into the `ResumeContent` schema
(a SINGLE `software_engineer` family), validate it, and write the YAML the
resume generator consumes.

Design notes:
- Claude extracts the structure (forced tool call), and we validate it
  against `ResumeContent` before anything is written.
- Honesty guardrails live in the prompt (prompts/resume_parse.txt): keep
  metrics verbatim, never invent experience. The parsed result is always
  shown to the user for review before any resume is generated.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Protocol

import yaml
from pydantic import BaseModel, Field

from ..llm import ClaudeClient
from .content_loader import ResumeContent

# Resume text is short; cap generously to keep token cost bounded and avoid
# pathological inputs.
RESUME_TEXT_TRUNCATE = 20000

# The single family this clone ships (see families.py).
FAMILY_NAME = "software_engineer"


# --- text extraction ----------------------------------------------------


def extract_resume_text(path: str | Path) -> str:
    """Extract plain text from a .pdf, .docx, or .txt/.md resume.

    Raises ValueError for unsupported extensions.
    """
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(p)
    if suffix == ".docx":
        return _extract_docx(p)
    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(
        f"Unsupported resume type {suffix!r}. Upload a PDF, DOCX, or TXT file."
    )


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = [para.text for para in doc.paragraphs]
    # Many resumes lay content out in tables — pull cell text too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(t for t in parts if t.strip()).strip()


# --- structured-output schema (lists, not dicts) ------------------------


class _ParsedBullet(BaseModel):
    text: str


class _ParsedExperience(BaseModel):
    company: str
    title: str
    location: str = ""
    start: str = ""
    end: str = ""
    bullets: list[_ParsedBullet] = Field(default_factory=list)


class _ParsedProject(BaseModel):
    title: str
    stack: str = ""
    bullets: list[_ParsedBullet] = Field(default_factory=list)


class _ParsedEducation(BaseModel):
    school: str
    degree: str
    location: str = ""
    start: str = ""
    end: str = ""
    honors: str = ""


class _ParsedSkills(BaseModel):
    languages: list[str] = Field(default_factory=list)
    data_platform: list[str] = Field(default_factory=list)
    cloud: list[str] = Field(default_factory=list)
    ml: list[str] = Field(default_factory=list)
    frontend: list[str] = Field(default_factory=list)
    other: list[str] = Field(default_factory=list)


class ParsedResume(BaseModel):
    """The flat shape the LLM returns. Converted to `ResumeContent` by
    `_parsed_to_content_dict`. Field names match the prompt exactly."""

    name: str
    email: str = ""
    phone: str = ""
    linkedin: str = ""
    location: str = ""
    summary: str = ""
    experiences: list[_ParsedExperience] = Field(default_factory=list)
    projects: list[_ParsedProject] = Field(default_factory=list)
    education: list[_ParsedEducation] = Field(default_factory=list)
    skills: _ParsedSkills = Field(default_factory=_ParsedSkills)
    languages_spoken: str = ""
    certifications: str = ""
    academic_background: str = ""


# --- LLM client ---------------------------------------------------------


class ResumeParseClient(Protocol):
    """Minimal interface for a resume-parsing LLM client.

    Production: `ClaudeResumeParseClient`. Tests: a fake returning a canned
    `ParsedResume`.
    """

    def parse(self, prompt: str) -> ParsedResume: ...


class ClaudeResumeParseClient:
    """Resume parsing via the shared ClaudeClient."""

    def __init__(self, *, llm: ClaudeClient, model: str) -> None:
        self._llm = llm
        self.model = model

    def parse(self, prompt: str) -> ParsedResume:
        return self._llm.structured(
            prompt,
            ParsedResume,
            model=self.model,
            max_tokens=8000,
            tool_name="submit_resume",
            tool_description="Submit the structured resume content.",
        )


# --- conversion + orchestration -----------------------------------------


def _slug(text: str, fallback: str) -> str:
    cleaned = "".join(c.lower() if c.isalnum() else "_" for c in text).strip("_")
    while "__" in cleaned:
        cleaned = cleaned.replace("__", "_")
    return cleaned or fallback


def _parsed_to_content_dict(parsed: ParsedResume) -> dict[str, Any]:
    """Convert the flat parse result into the ResumeContent dict shape.

    Assigns globally-unique, stable bullet IDs (b1, b2, … across all
    experiences; p1, p2, … across all projects) so the tailorer can merge
    reworded bullets back by id without collisions.
    """
    # Experiences → dict keyed by a slug of the company (deduped).
    experiences: dict[str, Any] = {}
    used_keys: set[str] = set()
    bullet_n = 0
    for i, exp in enumerate(parsed.experiences):
        key = _slug(exp.company, f"experience_{i + 1}")
        while key in used_keys:
            key = f"{key}_{i + 1}"
        used_keys.add(key)
        exp_bullets: list[dict[str, Any]] = []
        for b in exp.bullets:
            bullet_n += 1
            exp_bullets.append({"id": f"b{bullet_n}", "text": b.text, "tags": []})
        experiences[key] = {
            "company": exp.company,
            "title": exp.title,
            "location": exp.location,
            "start": exp.start,
            "end": exp.end,
            "bullets": exp_bullets,
        }

    # Projects → dict keyed by a slug of the title (deduped).
    projects: dict[str, Any] = {}
    used_proj_keys: set[str] = set()
    proj_n = 0
    for i, proj in enumerate(parsed.projects):
        key = _slug(proj.title, f"project_{i + 1}")
        while key in used_proj_keys:
            key = f"{key}_{i + 1}"
        used_proj_keys.add(key)
        proj_bullets: list[dict[str, Any]] = []
        for b in proj.bullets:
            proj_n += 1
            proj_bullets.append({"id": f"p{proj_n}", "text": b.text, "tags": []})
        projects[key] = {
            "title": proj.title,
            "stack": proj.stack,
            "bullets": proj_bullets,
        }

    return {
        "header": {
            "name": parsed.name,
            "email": parsed.email,
            "phone": parsed.phone,
            "linkedin": parsed.linkedin,
            "location": parsed.location,
        },
        "summaries": {FAMILY_NAME: parsed.summary},
        "experiences": experiences,
        "projects": projects,
        "education": [e.model_dump() for e in parsed.education],
        "skills": parsed.skills.model_dump(),
        "additional": {
            "languages": parsed.languages_spoken,
            "certifications": parsed.certifications,
            "academic_background": parsed.academic_background,
        },
    }


def parse_resume(
    text: str, client: ResumeParseClient, prompt_template: str
) -> ResumeContent:
    """Parse raw resume text into a validated `ResumeContent`.

    Uses `.replace` (not `.format`) to inject the resume text so literal
    braces in the resume don't break templating.
    """
    prompt = prompt_template.replace("{resume_text}", text[:RESUME_TEXT_TRUNCATE])
    parsed = client.parse(prompt)
    data = _parsed_to_content_dict(parsed)
    # Validate against the real schema — raises pydantic.ValidationError on
    # any mismatch, so we never write an unloadable file.
    return ResumeContent(**data)


_YAML_HEADER = """\
# Resume content — generated from your uploaded resume by the onboarding
# wizard. This is the single source of truth for every bullet, summary,
# experience, and education entry in a generated resume.
#
# Edit freely: fix any wording the parser got wrong, reorder bullets, drop
# ones you don't want. Keep the shape (see job_pipeline/resume/content_loader.py).
# Never inflate metrics — honest reframing only.
"""


def write_resume_content(content: ResumeContent, out_path: str | Path) -> None:
    """Write a validated ResumeContent to YAML with a schema header."""
    p = Path(out_path)
    body = yaml.safe_dump(
        content.model_dump(), sort_keys=False, allow_unicode=True, width=88
    )
    p.write_text(_YAML_HEADER + "\n" + body, encoding="utf-8")
